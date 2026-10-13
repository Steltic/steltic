"""The building template (.xlsx): write it (blank, or filled from a description) and read it back.

One tab per group. A blank cell is a missing value: the validator turns it into a question (key) or a proposed
default -- the reader never fills anything in. Plan positions take a grid label or a coordinate in feet.
"""
from __future__ import annotations

import hashlib
import pathlib

from . import schema as S

KV_PROJECT = [   # (row label, description path, note)
    ("Building name", "project/name", ""),
    ("Code edition", "project/code", "ASCE 7-22 (NL implements Chapter 16 of ASCE 7-22)"),
    ("SDS (g)", "seismic/SDS", ""), ("SD1 (g)", "seismic/SD1", ""), ("S1 (g)", "seismic/S1", "mapped"),
    ("TL (s)", "seismic/TL", "long-period transition"),
    ("Risk Category", "seismic/risk_category", "I, II, III or IV"), ("Ie", "seismic/Ie", "blank = Table 1.5-2 (proposed)"),
    ("Site class", "seismic/site_class", "A, B, BC, C, CD, D, DE, E"),
    ("Analysis procedure", "seismic/analysis", "ELF or MRSA (the linear procedure your design used)"),
    ("System X", "seismic/system/X", "SMF, IMF, OMF, SCBF, OCBF, BRBF, EBF, dual SMF+EBF / SMF+SCBF / SMF+BRBF / IMF+SCBF"),
    ("System Y", "seismic/system/Y", ""),
    ("R X", "seismic/R/X", "blank = Table 12.2-1 (proposed)"), ("R Y", "seismic/R/Y", ""),
    ("Cd X", "seismic/Cd/X", ""), ("Cd Y", "seismic/Cd/Y", ""), ("Omega0 X", "seismic/Om0/X", ""), ("Omega0 Y", "seismic/Om0/Y", ""),
    ("rho X", "seismic/rho/X", "redundancy factor used in your design"), ("rho Y", "seismic/rho/Y", ""),
    ("Diaphragm", "diaphragm", "rigid, flexible or semi-rigid (one for the building)"),
    ("Column bases (default)", "supports/default", "fixed or pinned; override per column on the Columns tab"),
    ("W-shape grade", "materials/W", "A992, A572 Gr 50, A36"), ("HSS grade", "materials/HSS", "A500 Gr C, A500 Gr B, A1085"),
    ("Existing structure", "materials/existing", "yes / no"), ("Existing: year of construction", "materials/existing_info/year", ""),
    ("Moment-frame beam bracing Lb (in)", "details/beam_bracing_in", "lateral bracing spacing of SMF / IMF beams (AISC 341 D1.2)"),
    ("BRB Fysc (ksi)", "materials/brb/Fysc_ksi", "BRBF only: adjusted / measured core yield stress"),
    ("BRB KF", "materials/brb/KF", "BRBF only: stiffness modification factor (supplier)"),
    ("BRB beta", "materials/brb/beta", "BRBF only: compression adjustment (supplier tests)"),
    ("BRB omega", "materials/brb/omega", "BRBF only: strain-hardening adjustment (supplier tests)"),
]
KV_LOADS = [
    ("Roof live Lr (psf)", "loads/Lr_psf", ""), ("Snow pf (psf)", "loads/snow_psf", "flat-roof snow"),
    ("Cladding (psf of wall)", "loads/clad_psf", "exterior wall weight"), ("Parapet height (ft)", "loads/parapet_ft", ""),
    ("Floor system", "loads/floor_system", "one-way or two-way"), ("Deck span", "loads/deck_span", "X or Y (one-way floors)"),
    ("Secondary framing spacing (in)", "loads/infill_spacing_in", "only when the secondary beams are NOT listed on the Beams tab"),
]
KV_WIND = [("Basic wind speed V (mph)", "wind/V_mph", ""), ("Exposure", "wind/exposure", "B, C or D"),
           ("Kzt", "wind/Kzt", ""), ("Enclosure", "wind/enclosure", "enclosed, partially enclosed, open")]
KV_RESULTS = [("T1 X (s)", "user_results/T1_s/X", "from your analysis"), ("T1 Y (s)", "user_results/T1_s/Y", ""),
              ("Base shear V X (kip)", "user_results/V_kip/X", "design base shear"), ("Base shear V Y (kip)", "user_results/V_kip/Y", ""),
              ("Seismic weight W (kip)", "user_results/W_kip", ""), ("Peak drift X (%)", "user_results/drift_pct/X", ""),
              ("Peak drift Y (%)", "user_results/drift_pct/Y", "")]
TABLES = {
    "Grid": ["Direction (X/Y)", "Label", "Coordinate (ft)"],
    "Levels": ["Name", "Elevation (ft)"],
    "Columns": ["Id", "X", "Y", "Base level", "Top level", "Section", "Strong axis (X/Y)", "Base fixity (optional)",
                "End release (none/base/top/both)"],
    "Beams": ["Id", "Level", "X i", "Y i", "X j", "Y j", "Section", "Release (none/I/J/both)", "EBF link (yes/no)",
              "RBS a (in)", "RBS b (in)", "RBS c (in)"],
    "Braces": ["Id", "X i", "Y i", "Level i", "X j", "Y j", "Level j", "Section", "BRB Asc (in2)", "System (optional)",
               "Buckling length factor (optional)"],
    "Level loads": ["Level (or floor / roof)", "D (psf)", "L (psf)", "Extra dead (psf)", "Seismic-only (psf)",
                    "Seismic weight (kip)", "Storage (yes/no)"],
    "Doublers": ["Level (or none)", "X", "Y", "Thickness (in)"],
    "Beam Lb": ["Section", "Lb (in)"],
    "Openings": ["Level", "X (lower-left grid corner)", "Y (lower-left grid corner)"],
}
README = [
    "Steltic building template -- a building YOU designed, for the nonlinear suite (pushover, NLRHA, DDM).",
    "HR models it exactly as given: nothing is resized, nothing is invented.",
    "",
    "Units: feet for grid coordinates and elevations; psf for area loads; kip; inches for detail dimensions; ksi.",
    "Plan positions (X, Y) take a grid label (e.g. A, 3) or a coordinate in feet.",
    "Levels: list from the base up; the first row is the base.",
    "Columns: one row per column stack or per segment; a column runs from its base level to its top level.",
    "Beams: one row per beam; a beam may span several bays (it is split where members frame in).",
    "  Release: none = moment connection at both ends; I or J = that end pinned; both = simple (shear) beam.",
    "  EBF link: mark the link segment (it must be its own row, between the brace work points).",
    "Braces: one row per diagonal; BRBs give Asc instead of a section.",
    "Floor plate: every bay whose four corner grid points are framed at a level carries floor; list atria and",
    "  shafts on the Openings tab (by the bay's lower-left grid corner).",
    "",
    "A BLANK cell is a missing value. HR then either asks you (key items) or proposes a default that you must",
    "accept before anything is built -- every accepted default is printed in every report.",
    "Version 1 systems: SMF, IMF, OMF, SCBF, OCBF, BRBF, EBF and the dual combinations listed on Project.",
    "Not supported (refused, never simplified): plate shear walls (SPSW, C-PSW), STMF, partially restrained",
    "connections, sloped or plan-diagonal beams, more than one diaphragm type, per-member steel grades.",
]


def _v(x):
    if x is None:
        return None
    if isinstance(x, str):
        s = x.strip()
        if s == "":
            return None
        try:
            return float(s) if any(c.isdigit() for c in s) and s.replace(".", "", 1).replace("-", "", 1).isdigit() else s
        except ValueError:
            return s
    if isinstance(x, float) and x.is_integer():
        return x
    return x


def _pt(x, y, level=None):
    p = {"x": _label_or_num(x), "y": _label_or_num(y)}
    if level is not None:
        p["level"] = str(_v(level)) if _v(level) is not None else None
    return p


def _label_or_num(v):
    v = _v(v)
    if isinstance(v, float):
        return v
    return None if v is None else str(v)


def _yes(v):
    return str(v or "").strip().lower() in ("yes", "y", "true", "1", "x")


def write_template(path, desc=None):
    """Write the template; filled from `desc` when given (used for examples and round-trip tests)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    d = desc or {}
    wb = Workbook()
    ws = wb.active; ws.title = "Read me"
    for r, line in enumerate(README, 1):
        ws.cell(r, 1, line)
    ws.column_dimensions["A"].width = 110
    bold = Font(bold=True); head = PatternFill("solid", fgColor="DDE6F0")

    def kv_sheet(title, rows):
        s = wb.create_sheet(title)
        s.append(["Item", "Value", "Note"])
        for c in s[1]:
            c.font = bold; c.fill = head
        for lab, path, note in rows:
            val = S.get_path(d, path)
            if path == "materials/existing" and val is not None:
                val = "yes" if val else "no"
            s.append([lab, val, note])
        s.column_dimensions["A"].width = 32; s.column_dimensions["B"].width = 22; s.column_dimensions["C"].width = 90
        return s

    kv_sheet("Project", KV_PROJECT)
    kv_sheet("Loads", KV_LOADS)
    kv_sheet("Wind", KV_WIND)
    kv_sheet("Your results", KV_RESULTS)
    sheets = {}
    for t, hdr in TABLES.items():
        s = wb.create_sheet(t)
        s.append(hdr)
        for c in s[1]:
            c.font = bold; c.fill = head; c.alignment = Alignment(wrap_text=True)
        for i in range(len(hdr)):
            s.column_dimensions[chr(65 + i)].width = 16
        sheets[t] = s
    for ax in ("x", "y"):
        for g in ((d.get("grid") or {}).get(ax) or []):
            sheets["Grid"].append([ax.upper(), g.get("label"), g.get("coord_ft")])
    for lv in d.get("levels") or []:
        sheets["Levels"].append([lv.get("name"), lv.get("elev_ft")])
    sup_at = {(str(a.get("x")), str(a.get("y"))): a.get("fixity") for a in ((d.get("supports") or {}).get("at") or [])}
    for c in d.get("columns") or []:
        sheets["Columns"].append([c.get("id"), c.get("x"), c.get("y"), c.get("base_level"), c.get("top_level"), c.get("section"),
                                  c.get("strong_axis"), sup_at.get((str(c.get("x")), str(c.get("y")))), c.get("release")])
    for b in d.get("beams") or []:
        rb = b.get("rbs") if isinstance(b.get("rbs"), dict) else {}
        sheets["Beams"].append([b.get("id"), b.get("level"), (b.get("i") or {}).get("x"), (b.get("i") or {}).get("y"),
                                (b.get("j") or {}).get("x"), (b.get("j") or {}).get("y"), b.get("section"), b.get("release"),
                                "yes" if b.get("link") else None, rb.get("a_in"), rb.get("b_in"), rb.get("c_in")])
    for b in d.get("braces") or []:
        i, j = b.get("i") or {}, b.get("j") or {}
        sheets["Braces"].append([b.get("id"), i.get("x"), i.get("y"), i.get("level"), j.get("x"), j.get("y"), j.get("level"),
                                 b.get("section"), (b.get("brb") or {}).get("Asc_in2"), b.get("system"), b.get("length_factor")])
    lo = d.get("loads") or {}
    keys = []
    for k in ("D_psf", "L_psf", "extra_dead_psf", "seismic_only_psf", "seismic_weight_kip"):
        v = lo.get(k)
        if isinstance(v, dict):
            keys += [n for n in v if n not in keys]
        elif v is not None and "all" not in keys:
            keys.append("all")
    stor = set(map(str, lo.get("storage_levels") or []))
    for n in keys + [n for n in stor if n not in keys]:
        row = [n]
        for k in ("D_psf", "L_psf", "extra_dead_psf", "seismic_only_psf", "seismic_weight_kip"):
            v = lo.get(k)
            row.append(v.get(n) if isinstance(v, dict) else (v if n == "all" else None))
        row.append("yes" if n in stor else None)
        sheets["Level loads"].append(row)
    de = d.get("details") or {}
    if de.get("doublers") == []:
        sheets["Doublers"].append(["none"])
    for x in de.get("doublers") or []:
        sheets["Doublers"].append([x.get("level"), x.get("x"), x.get("y"), x.get("t_in")])
    for sec, lb in (de.get("beam_Lb_in") or {}).items():
        sheets["Beam Lb"].append([sec, lb])
    for o in d.get("openings") or []:
        sheets["Openings"].append([o.get("level"), o.get("x"), o.get("y")])
    # dropdowns
    def dv(sheet, col, opts):
        v = DataValidation(type="list", formula1='"%s"' % ",".join(opts), allow_blank=True)
        sheets[sheet].add_data_validation(v); v.add("%s2:%s5000" % (col, col))
    dv("Grid", "A", ["X", "Y"]); dv("Columns", "G", ["X", "Y"]); dv("Columns", "H", ["fixed", "pinned"])
    dv("Columns", "I", ["none", "base", "top", "both"])
    dv("Beams", "H", ["none", "I", "J", "both"]); dv("Beams", "I", ["yes", "no"])
    dv("Braces", "J", ["SCBF", "OCBF", "BRBF", "EBF"])
    path = pathlib.Path(path)
    wb.save(path)
    return path


def _rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    hdr = [str(h or "").strip().lower() for h in rows[0]]
    return hdr, [r for r in rows[1:] if any(_v(c) is not None for c in r)]


def _col(hdr, *names):
    for n in names:
        for i, h in enumerate(hdr):
            if h.startswith(n):
                return i
    return None


def read_template(path) -> dict:
    from openpyxl import load_workbook
    path = pathlib.Path(path)
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    wb = load_workbook(path, data_only=True)
    d = S.new_description("template", path.name, sha)
    names = {ws.title.strip().lower(): ws for ws in wb.worksheets}

    def kv(title, spec):
        ws = names.get(title.lower())
        if ws is None:
            return
        lab2path = {lab.lower(): p for lab, p, _ in spec}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or r[0] is None:
                continue
            p = lab2path.get(str(r[0]).strip().lower())
            val = _v(r[1] if len(r) > 1 else None)
            if p and val is not None:
                if p == "materials/existing":
                    val = _yes(val)
                elif p in ("seismic/risk_category", "seismic/site_class", "materials/existing_info/year", "project/name") and isinstance(val, float):
                    val = str(int(val)) if val.is_integer() else str(val)
                S.set_path(d, p, val)
                S.mark(d, p, "file")

    kv("Project", KV_PROJECT)
    kv("Loads", KV_LOADS)
    kv("Wind", KV_WIND)
    kv("Your results", KV_RESULTS)
    if not d.get("wind") or not isinstance(d.get("wind"), dict) or not d["wind"].get("V_mph"):
        d["wind"] = None
    if not d.get("user_results") or not any(v for v in (d.get("user_results") or {}).values()):
        d["user_results"] = None

    ws = names.get("grid")
    if ws:
        hdr, rows = _rows(ws)
        for r in rows:
            ax = str(_v(r[0]) or "").lower()
            if ax in ("x", "y"):
                d["grid"][ax].append({"label": _label_or_num(r[1]) if not isinstance(_v(r[1]), float) else ("%g" % _v(r[1])),
                                      "coord_ft": _v(r[2])})
        S.mark(d, "grid", "file")
    ws = names.get("levels")
    if ws:
        hdr, rows = _rows(ws)
        for r in rows:
            n = _v(r[0])
            d["levels"].append({"name": ("%g" % n) if isinstance(n, float) else n, "elev_ft": _v(r[1])})
        S.mark(d, "levels", "file")

    def lvl(v):
        v = _v(v)
        return ("%g" % v) if isinstance(v, float) else v

    ws = names.get("columns")
    sup_at = []
    if ws:
        hdr, rows = _rows(ws)
        for n, r in enumerate(rows):
            g = lambda *k: _v(r[_col(hdr, *k)]) if _col(hdr, *k) is not None and _col(hdr, *k) < len(r) else None
            c = {"id": str(g("id") or "C%d" % (n + 1)), "x": _label_or_num(g("x")), "y": _label_or_num(g("y")),
                 "base_level": lvl(g("base level")), "top_level": lvl(g("top level")), "section": g("section"),
                 "strong_axis": (str(g("strong")).upper() if g("strong") else None), "_where": "Columns row %d" % (n + 2)}
            if g("end release"):
                c["release"] = str(g("end release")).strip().lower()
            d["columns"].append(c)
            if g("base fixity"):
                sup_at.append({"x": c["x"], "y": c["y"], "fixity": str(g("base fixity")).lower()})
    if sup_at:
        d.setdefault("supports", {})["at"] = sup_at
    ws = names.get("beams")
    if ws:
        hdr, rows = _rows(ws)
        for n, r in enumerate(rows):
            g = lambda *k: _v(r[_col(hdr, *k)]) if _col(hdr, *k) is not None and _col(hdr, *k) < len(r) else None
            b = {"id": str(g("id") or "B%d" % (n + 1)), "level": lvl(g("level")), "i": _pt(g("x i"), g("y i")),
                 "j": _pt(g("x j"), g("y j")), "section": g("section"),
                 "release": (str(g("release")).strip() if g("release") is not None else None),
                 "link": _yes(g("ebf link")), "_where": "Beams row %d" % (n + 2)}
            ra, rb_, rc = g("rbs a"), g("rbs b"), g("rbs c")
            if any(x is not None for x in (ra, rb_, rc)):
                b["rbs"] = {"a_in": ra, "b_in": rb_, "c_in": rc}
            d["beams"].append(b)
    ws = names.get("braces")
    if ws:
        hdr, rows = _rows(ws)
        for n, r in enumerate(rows):
            g = lambda *k: _v(r[_col(hdr, *k)]) if _col(hdr, *k) is not None and _col(hdr, *k) < len(r) else None
            b = {"id": str(g("id") or "D%d" % (n + 1)), "i": _pt(g("x i"), g("y i"), g("level i")),
                 "j": _pt(g("x j"), g("y j"), g("level j")), "section": g("section"), "_where": "Braces row %d" % (n + 2)}
            if g("brb asc") is not None:
                b["brb"] = {"Asc_in2": g("brb asc")}
            if g("system"):
                b["system"] = str(g("system"))
            if g("buckling length") is not None:
                b["length_factor"] = g("buckling length")
            d["braces"].append(b)
    ws = names.get("level loads")
    if ws:
        hdr, rows = _rows(ws)
        lo = d.setdefault("loads", {})
        stor = []
        for r in rows:
            n = lvl(r[0])
            for ci, k in ((1, "D_psf"), (2, "L_psf"), (3, "extra_dead_psf"), (4, "seismic_only_psf"), (5, "seismic_weight_kip")):
                v = _v(r[ci]) if ci < len(r) else None
                if v is None:
                    continue
                if str(n).lower() == "all":
                    lo[k] = v
                else:
                    lo.setdefault(k, {}) if not isinstance(lo.get(k), (int, float)) else None
                    if isinstance(lo.get(k), dict):
                        lo[k][str(n).lower() if str(n).lower() in ("floor", "roof") else str(n)] = v
            if len(r) > 6 and _yes(r[6]):
                stor.append(str(n))
        if stor:
            lo["storage_levels"] = stor
        S.mark(d, "loads", "file")
    ws = names.get("doublers")
    if ws:
        hdr, rows = _rows(ws)
        if any(str(_v(r[0]) or "").lower() == "none" for r in rows):
            d.setdefault("details", {})["doublers"] = "none"
        else:
            dl = [{"level": lvl(r[0]), "x": _label_or_num(r[1]), "y": _label_or_num(r[2]), "t_in": _v(r[3] if len(r) > 3 else None)} for r in rows]
            d.setdefault("details", {})["doublers"] = dl if dl else None
    ws = names.get("openings")
    if ws:
        hdr, rows = _rows(ws)
        if rows:
            d["openings"] = [{"level": lvl(r[0]), "x": _label_or_num(r[1]), "y": _label_or_num(r[2])} for r in rows]
    ws = names.get("beam lb")
    if ws:
        hdr, rows = _rows(ws)
        if rows:
            d.setdefault("details", {})["beam_Lb_in"] = {str(_v(r[0])): _v(r[1]) for r in rows if _v(r[1]) is not None}
    # a doubler tab that exists but is empty means "not answered" (None), not "none"
    if (d.get("details") or {}).get("doublers") == []:
        d["details"]["doublers"] = None
    if (d.get("details") or {}).get("doublers") == "none":
        d["details"]["doublers"] = []
    S.mark(d, "columns", "file"); S.mark(d, "beams", "file"); S.mark(d, "braces", "file")
    return d
